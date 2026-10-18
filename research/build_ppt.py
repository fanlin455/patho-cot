# -*- coding: utf-8 -*-
"""生成《智析病理》作品答辩 PPT（16:9）"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

# 路径无关：图片素材取脚本同目录的 demo_assets，产出放到上一级的「答辩材料」
_HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(_HERE, 'demo_assets')
OUT_DIR = os.path.join(os.path.dirname(_HERE), '答辩材料')
os.makedirs(OUT_DIR, exist_ok=True)
PPTX = os.path.join(OUT_DIR, '智析病理-作品答辩.pptx')

DARK = RGBColor(0x0B, 0x0F, 0x17)
BLUE = RGBColor(0x25, 0x63, 0xEB)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
TEAL = RGBColor(0x0D, 0x94, 0x88)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
PURPLE = RGBColor(0x93, 0x33, 0xEA)
GREY = RGBColor(0x64, 0x74, 0x8B)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'Microsoft YaHei'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def tb(slide, x, y, w, h, text, size=16, color=DARK, bold=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, line=1.35):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split('\n')
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
    return box


def rect(slide, x, y, w, h, color, line=False):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if line:
        sh.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = 0.06
    except Exception:
        pass
    return sh


def header(slide, title, sub='', accent=BLUE):
    rect(slide, 0, 0, 0.16, 7.5, accent)
    tb(slide, 0.62, 0.36, 11.8, 0.7, title, size=26, bold=True)
    if sub:
        tb(slide, 0.66, 1.02, 11.8, 0.4, sub, size=12.5, color=GREY)


def pic(slide, path, x, y, w=None, h=None):
    p = os.path.join(ASSETS, path)
    kw = {}
    if w:
        kw['width'] = Inches(w)
    if h:
        kw['height'] = Inches(h)
    return slide.shapes.add_picture(p, Inches(x), Inches(y), **kw)


def kpi(slide, x, y, w, value, label, color):
    rect(slide, x, y, w, 1.28, LIGHT)
    tb(slide, x, y + 0.14, w, 0.55, value, size=25, bold=True, color=color, align=PP_ALIGN.CENTER)
    tb(slide, x, y + 0.76, w, 0.36, label, size=11.5, color=GREY, align=PP_ALIGN.CENTER)


def bullets(slide, x, y, w, items, size=14, gap=0.52):
    for i, (t, c) in enumerate(items):
        tb(slide, x, y + i * gap, 0.28, 0.4, '●', size=size - 3, color=c)
        tb(slide, x + 0.32, y + i * gap, w - 0.32, 0.5, t, size=size, color=RGBColor(0x33, 0x41, 0x55))


# ============ 1 封面 ============
s = prs.slides.add_slide(BLANK)
s.shapes.add_picture(os.path.join(ASSETS, 'card_title.png'), 0, 0, width=prs.slide_width)
tb(s, 0.9, 6.62, 11.5, 0.4, '数字媒体竞赛参赛作品 · 作品答辩', size=13, color=GREY)

# ============ 2 作品简介 ============
s = prs.slides.add_slide(BLANK)
header(s, '一、作品简介与设计出发点', 'Why we built it')
tb(s, 0.66, 1.55, 6.0, 0.5, '问题：病理诊断存在能力门槛', size=17, bold=True, color=BLUE)
bullets(s, 0.66, 2.15, 6.1, [
    ('基层医院与教学场景缺乏病理读片能力', GREY),
    ('病理医生培养周期长，切片判读高度依赖经验', GREY),
    ('现有 AI 方案多为"黑箱给结论"，医生无法复核推理过程', GREY),
], size=13.5)
tb(s, 0.66, 4.25, 6.0, 0.5, '我们的思路：把推理链本身可视化', size=17, bold=True, color=TEAL)
bullets(s, 0.66, 4.85, 6.1, [
    ('把"看片 → 定位 → 定性 → 出报告"拆成六个可回看、可单步重试的环节', GREY),
    ('每一级都给出量化证据（分数 / 计数 / 概率）而非结论黑箱', GREY),
    ('定位是"辅助与教学"，不替代医生诊断', GREY),
], size=13.5)
kpi(s, 7.15, 1.7, 2.55, '6 级', '引导式分析步骤', BLUE)
kpi(s, 9.95, 1.7, 2.55, '19 类', '器官来源识别', INDIGO)
kpi(s, 7.15, 3.25, 2.55, '23 例', '已入库分析记录', TEAL)
kpi(s, 9.95, 3.25, 2.55, '99.75%', '器官分类准确率', RED)
rect(s, 7.15, 4.95, 5.35, 1.85, LIGHT)
tb(s, 7.45, 5.15, 4.8, 1.5,
   '全流程本地运行\n不依赖云端 API，病例数据不出本机\n（AI 报告由内置本地小模型生成）',
   size=12.5, color=RGBColor(0x33, 0x41, 0x55), line=1.5)

# ============ 3 架构 ============
s = prs.slides.add_slide(BLANK)
header(s, '二、系统总体架构', '六级流水线 · 统一推理引擎')
pic(s, 'diagram_arch.png', 0.55, 1.6, w=12.2)
tb(s, 0.66, 6.55, 12.0, 0.5,
   '技术栈：Python 3.12 · PyTorch (CUDA) · YOLOv8 · Phikon-v2(DINOv2) · diffusers · FastAPI + Jinja2 · 内置本地 LLM（可选升级 Ollama）',
   size=11.5, color=GREY)

# ============ 4 界面设计 ============
s = prs.slides.add_slide(BLANK)
header(s, '三、界面设计', '引导式步骤向导 · 浅色玻璃拟态', TEAL)
pic(s, '01_home.png', 0.6, 1.5, w=7.7)
tb(s, 8.65, 1.6, 4.1, 0.4, '设计要点', size=16, bold=True, color=TEAL)
bullets(s, 8.65, 2.15, 4.15, [
    ('顶部步骤条实时显示进度与状态', GREY),
    ('六个步骤各有主色，贯穿徽章/图表/卡片光晕', GREY),
    ('卡片光晕跟随鼠标，颜色取自卡片主色', GREY),
    ('导航与按钮统一为黑色胶囊 + 粘液融合动效', GREY),
    ('支持打印/存 PDF，打印时自动隐藏导航', GREY),
], size=12.5, gap=0.62)

# ============ 5 创新点一 ============
s = prs.slides.add_slide(BLANK)
header(s, '四、创新点①：冻结病理基础模型替代自训练小模型', '器官分类 · 零训练 +5.57pp', BLUE)
pic(s, 'chart_organ.png', 0.55, 1.5, w=7.5)
tb(s, 8.35, 1.6, 4.4, 0.4, '做法', size=15, bold=True, color=BLUE)
bullets(s, 8.35, 2.12, 4.45, [
    ('用病理专用基础模型 Phikon-v2（DINOv2 ViT-L/16）提特征', GREY),
    ('只训练一个线性分类头，骨干完全冻结', GREY),
    ('kNN 只需特征库，连训练都不需要', GREY),
], size=12.5, gap=0.66)
rect(s, 8.35, 4.35, 4.4, 2.15, LIGHT)
tb(s, 8.6, 4.55, 3.95, 1.8,
   '为什么重要\n· 94.18% → 99.75%，17/19 类完全正确\n· 特征抽取 7901 张仅需 103 秒\n· 零训练成本，换数据集即可复用',
   size=12, color=RGBColor(0x33, 0x41, 0x55), line=1.5)

# ============ 6 泄漏验证 ============
s = prs.slides.add_slide(BLANK)
header(s, '五、创新点①验证：数据泄漏三重排查', '高准确率必须经得起质疑', TEAL)
tb(s, 0.66, 1.5, 12.0, 0.6,
   '数据集源自 TCGA，而 Phikon-v2 预训练也包含 TCGA —— 存在"见过考题"的可能，因此必须独立验证。',
   size=13.5, color=RGBColor(0x33, 0x41, 0x55))
rows = [
    ('训练/测试近邻重复数', '0', '无重复图'),
    ('最近邻距离（中位 / 最小）', '0.136 / 0.030', '无过近样本'),
    ('分类间隔 margin（中位）', '0.253', '间隔健康，非死记硬背'),
    ('NCT-CRC 独立队列（非 TCGA，1800 张）', '99.78%', '跨域泛化确认'),
]
y = 2.45
for i, (a, b, c) in enumerate(rows):
    bg = LIGHT if i % 2 == 0 else WHITE
    rect(s, 0.66, y + i * 0.76, 12.0, 0.68, bg)
    tb(s, 0.9, y + i * 0.76 + 0.13, 5.4, 0.42, a, size=13.5, color=RGBColor(0x33, 0x41, 0x55))
    tb(s, 6.4, y + i * 0.76 + 0.13, 2.4, 0.42, b, size=13.5, bold=True,
       color=TEAL if i == 3 else DARK, align=PP_ALIGN.CENTER)
    tb(s, 8.9, y + i * 0.76 + 0.13, 3.6, 0.42, c, size=12.5, color=GREY)
tb(s, 0.66, 5.75, 12.0, 0.9,
   '结论：99.75% 来自真实泛化能力，而非数据泄漏 —— 因此可以放心接入生产流程。',
   size=15, bold=True, color=TEAL)

# ============ 7 创新点二 ============
s = prs.slides.add_slide(BLANK)
header(s, '六、创新点②：把"循环定义"的任务重构掉', '病变性质 · AUC 87.4% → 96.2%', PURPLE)
pic(s, 'chart_nature.png', 0.55, 1.5, w=7.0)
tb(s, 7.85, 1.55, 4.9, 0.4, '实验中发现的关键问题', size=15, bold=True, color=RED)
rect(s, 7.85, 2.05, 4.9, 1.65, RGBColor(0xFE, 0xF2, 0xF2))
tb(s, 8.1, 2.22, 4.45, 1.4,
   '标签本身就是"病变细胞占比 ≥ 0.4936 → 异常"，\n而这个占比在部署管线里可直接算出。\n即：模型在用几何特征反推一个能直接计算的阈值。',
   size=11.5, color=RGBColor(0x7F, 0x1D, 0x1D), line=1.45)
tb(s, 7.85, 3.9, 4.9, 0.4, '重构方案：确定性规则 + 图像二次判读', size=14.5, bold=True, color=TEAL)
bullets(s, 7.85, 4.4, 4.9, [
    ('主判据：ratio 阈值规则（确定性命中）', GREY),
    ('二次判读：冻结 Phikon-v2 线性头，AUC 96.2%', GREY),
    ('图像细胞中位数仅 20 个 → 弱标签噪声极大，已在报告中说明', GREY),
], size=12, gap=0.66)

# ============ 8 创新点三 ============
s = prs.slides.add_slide(BLANK)
header(s, '七、创新点③：生成式模型做无监督异常检测', '扩散重建误差 + 异常热力图（实验性）', AMBER)
pic(s, 'chart_p2.png', 0.5, 1.45, w=8.4)
pic(s, '05_anomaly_heatmap.png', 9.15, 1.55, w=3.85)
tb(s, 9.15, 3.75, 3.9, 0.4, '异常热力图（左：原图 中：热力 右：叠加）', size=10.5, color=GREY)
bullets(s, 9.15, 4.3, 3.9, [
    ('仅用"正常"样本训练 DDPM', GREY),
    ('重建误差 = 异常分数', GREY),
    ('接入阶段 4，附带输出分数 + 热力图', GREY),
], size=11.5, gap=0.58)
rect(s, 0.6, 5.9, 12.2, 0.95, RGBColor(0xFF, 0xFB, 0xEB))
tb(s, 0.9, 6.08, 11.6, 0.7,
   '诚实结论：AUROC 0.67 优于随机但远未达可用水平（受限于单机算力与弱标签噪声）。'
   '作品明确标注为"实验性参考"，界面与 AI 报告中均不作为诊断结论，定位为技术探索与对比实验。',
   size=12, color=RGBColor(0x92, 0x40, 0x0E), line=1.45)

# ============ 9 交互创新 ============
s = prs.slides.add_slide(BLANK)
header(s, '八、界面交互创新', '动效为功能服务，且保持克制', INDIGO)
bullets(s, 0.66, 1.6, 6.2, [
    ('GooeyNav 粘液融合导航：点击时黑色胶囊带黏液团在按钮间融合移动', GREY),
    ('为适配浅色主题做了完整配色镜像（官方为深底白胶囊 + lighten）', GREY),
    ('以 alpha 通道 goo 滤镜替代官方 blur+contrast，解决溢出与边缘失真', GREY),
    ('动效时长压缩至约 0.6 秒，只在点击时触发，不影响操作效率', GREY),
    ('SpotlightCard 光晕颜色取自各卡片主色，强化六步色彩区分', GREY),
], size=13, gap=0.66)
pic(s, '04_multi.png', 7.1, 1.6, w=5.75)
tb(s, 0.66, 5.35, 6.2, 1.5,
   '设计原则：动效必须"看得出、不碍事、可关闭"。\n所有导航与按钮统一为同一套视觉语言，降低学习成本；\n步骤标识保持克制，不做多余动效。',
   size=12.5, color=RGBColor(0x33, 0x41, 0x55), line=1.55)

# ============ 10 完整链路实测 ============
s = prs.slides.add_slide(BLANK)
header(s, '九、完整分析链路实测', '真实病例（多 patch 批量分析，记录 #22）', TEAL)
pic(s, '03_detail.png', 0.6, 1.42, h=5.75)
tb(s, 6.05, 1.55, 6.75, 0.4, '一次分析产出的量化证据', size=16, bold=True, color=TEAL)
rows = [
    ('多 patch 汇总', '5 个 patch，3 个合格，病变细胞合计 313'),
    ('Step 1 质量检测', '可用性 0.7547 · 对焦 10.0 · 染色 4.0 · 组织占比 26.5%'),
    ('Step 2 器官分类', '头颈部（置信度 89.68%）'),
    ('Step 3 细胞检测', '检出病变细胞 201 / 非病变 0 / 合计 201'),
    ('Step 4 病变性质', '异常（置信度 99.0%）'),
    ('Step 5 组织亚型', '肿瘤相关间质（28.82%）'),
]
for i, (a, b) in enumerate(rows):
    yy = 2.15 + i * 0.85
    rect(s, 6.05, yy, 6.75, 0.75, LIGHT if i % 2 == 0 else WHITE)
    tb(s, 6.28, yy + 0.08, 2.0, 0.35, a, size=12, bold=True, color=DARK)
    tb(s, 6.28, yy + 0.4, 6.3, 0.32, b, size=11.5, color=GREY)

# ============ 11 AI 报告 ============
s = prs.slides.add_slide(BLANK)
header(s, '十、AI 结构化报告生成', '本地部署 LLM，病例数据不出本机', INDIGO)
if os.path.exists(os.path.join(ASSETS, '06_report.png')):
    pic(s, '06_report.png', 0.6, 1.5, w=7.9)
    tb(s, 8.75, 1.6, 4.0, 0.4, '报告生成链路', size=15, bold=True, color=INDIGO)
    bullets(s, 8.75, 2.15, 4.05, [
        ('汇总前五步结构化结果 + 病情描述', GREY),
        ('内置本地小模型生成中文报告，无需联网、无需额外服务', GREY),
        ('若检测到本地 Ollama，自动升级为更大参数模型', GREY),
        ('自动包含免责声明，统一 Normal/Abnormal 口径', GREY),
        ('支持复制与打印/存 PDF', GREY),
    ], size=12.5, gap=0.68)
    rect(s, 8.75, 5.05, 4.05, 1.6, LIGHT)
    tb(s, 9.0, 5.22, 3.55, 1.3,
       '工程说明\n病例数据全程不出本机；显存不足时会主动通知\nOllama 卸载模型让出显存，保证流水线稳定。',
       size=11.5, color=RGBColor(0x33, 0x41, 0x55), line=1.5)
else:
    tb(s, 0.66, 3.0, 12.0, 1.0, '（待补充真实报告截图）', size=16, color=GREY)

# ============ 12 工程亮点 ============
s = prs.slides.add_slide(BLANK)
header(s, '十一、工程实现亮点', '让系统真正跑得起来', BLUE)
cards = [
    ('显存调度', '8GB 显存下单模型驻留；步骤间释放；\n推理前通知 Ollama 让出显存', BLUE),
    ('路径无关', '所有路径由项目根目录推导，\n整个文件夹可任意搬运/换机', TEAL),
    ('步骤级容错', '质量不合格可人工确认后继续；\n任一步失败可单独重试', AMBER),
    ('配置驱动', 'config.yaml 统一切换新旧模型、\n阈值与开关，改配置即生效', INDIGO),
]
for i, (t, d, c) in enumerate(cards):
    x = 0.66 + (i % 2) * 6.25
    y = 1.7 + (i // 2) * 2.4
    rect(s, x, y, 5.9, 2.1, LIGHT)
    rect(s, x, y, 0.12, 2.1, c)
    tb(s, x + 0.42, y + 0.28, 5.2, 0.45, t, size=17, bold=True, color=c)
    tb(s, x + 0.42, y + 0.92, 5.2, 1.0, d, size=12.5, color=RGBColor(0x33, 0x41, 0x55), line=1.45)

# ============ 13 结果汇总 ============
s = prs.slides.add_slide(BLANK)
header(s, '十二、实验结果汇总', '全部指标均来自本机实测', TEAL)
hdr = ['升级项', '原方案', '新方案', '提升', '状态']
colw = [3.4, 2.3, 2.6, 1.7, 2.0]
data = [
    ('器官分类', 'ViT-B/16  94.18%', 'Phikon-v2+kNN  99.75%', '+5.57pp', '已上线'),
    ('病变性质', 'GCN  AUC 87.4%', '规则+Phikon  AUC 96.2%', '+8.8pp', '已上线'),
    ('异常检测', '无', 'AnoDDPM  AUROC 0.67', '新增能力', '实验性'),
]
x0 = 0.66
rect(s, x0, 1.65, sum(colw), 0.62, DARK)
cx = x0
for i, h in enumerate(hdr):
    tb(s, cx + 0.18, 1.76, colw[i] - 0.2, 0.4, h, size=12.5, bold=True, color=WHITE)
    cx += colw[i]
for r, row in enumerate(data):
    y = 2.27 + r * 0.78
    rect(s, x0, y, sum(colw), 0.72, LIGHT if r % 2 == 0 else WHITE)
    cx = x0
    for i, cell in enumerate(row):
        col = TEAL if i == 3 else (GREY if i == 1 else RGBColor(0x33, 0x41, 0x55))
        bold = (i == 3)
        tb(s, cx + 0.18, y + 0.14, colw[i] - 0.2, 0.45, cell, size=12, color=col, bold=bold)
        cx += colw[i]
tb(s, 0.66, 4.95, 12.0, 0.45, '补充说明', size=15, bold=True, color=BLUE)
bullets(s, 0.66, 5.5, 12.0, [
    ('器官分类：NCT-CRC 独立队列（非 TCGA）复测 99.78%，排除数据泄漏', GREY),
    ('病变性质：原任务标签为"循环定义"，本次先修正任务定义再换模型', GREY),
    ('异常检测：AUROC 0.67 未达可用水平，明确标注为实验性，不参与诊断结论', GREY),
], size=12.5, gap=0.52)

# ============ 14 限制与后续 ============
s = prs.slides.add_slide(BLANK)
header(s, '十三、已知限制与后续工作', '如实说明', RED)
tb(s, 0.66, 1.55, 6.0, 0.4, '已知限制', size=17, bold=True, color=RED)
bullets(s, 0.66, 2.15, 6.1, [
    ('异常检测 AUROC 0.67，仅作技术探索展示', GREY),
    ('报告由内置本地小模型生成，无需联网；部署 Ollama 可自动升级', GREY),
    ('病变性质标签为弱标签（由细胞占比规则生成），存在固有噪声', GREY),
    ('系统定位为科研/教学演示，不可用于临床诊断', GREY),
    ('8GB 显存下 SAM 像素级分割默认关闭', GREY),
], size=13, gap=0.72)
tb(s, 7.15, 1.55, 5.6, 0.4, '后续工作', size=17, bold=True, color=TEAL)
bullets(s, 7.15, 2.15, 5.6, [
    ('引入真实病理标注，替代弱标签重训病变性质', GREY),
    ('更大算力下提升扩散异常检测分辨率与训练步数', GREY),
    ('全切片（WSI）金字塔推理与多尺度融合', GREY),
    ('报告结构化输出与诊断编码标准化', GREY),
], size=13, gap=0.72)

# ============ 15 结束 ============
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, DARK)
tb(s, 1.2, 2.55, 11.0, 0.9, '智析病理  Patho-COT', size=40, bold=True, color=WHITE)
tb(s, 1.25, 3.65, 11.0, 0.5, '把病理诊断的推理链，变成看得见、可复核的六步', size=16, color=RGBColor(0x94, 0xA3, 0xB8))
rect(s, 1.25, 4.45, 2.2, 0.05, BLUE)
tb(s, 1.25, 4.85, 11.0, 0.5, '感谢观看 · 恳请批评指正', size=15, color=RGBColor(0xCB, 0xD5, 0xE1))
tb(s, 1.25, 6.5, 11.0, 0.4,
   '本作品为科研与教学演示系统，所有输出仅供参考，不可作为临床诊断依据。',
   size=10.5, color=RGBColor(0x64, 0x74, 0x8B))

prs.save(PPTX)
print('SLIDES:', len(prs.slides.__iter__.__self__._sldIdLst))
print('SAVED:', PPTX)
