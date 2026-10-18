# -*- coding: utf-8 -*-
"""读回 PPTX，核实幻灯片尺寸与形状坐标"""
from pptx import Presentation
from pptx.util import Emu

P = r'C:\Desktop\数媒比赛-上交源文件\比赛材料\智析病理-作品答辩.pptx'
prs = Presentation(P)
print('slide_width  = %.3f in' % (prs.slide_width / 914400))
print('slide_height = %.3f in' % (prs.slide_height / 914400))
print('pages =', len(prs.slides))

for idx in (1, 2, 3):
    s = prs.slides[idx]
    print('\n--- page %d ---' % (idx + 1))
    for sh in s.shapes:
        try:
            l = sh.left / 914400
            t = sh.top / 914400
            w = sh.width / 914400
        except Exception:
            l = t = w = None
        txt = ''
        if sh.has_text_frame:
            txt = sh.text_frame.text.replace('\n', ' / ')[:26]
        print('  %-14s L=%6.2f  W=%6.2f  R=%6.2f   %s' % (
            sh.shape_type, l if l is not None else -1, w if w is not None else -1,
            (l + w) if l is not None else -1, txt))
