# -*- coding: utf-8 -*-
"""标尺测试：在已知位置放色块，检查渲染是否与文件坐标一致"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

P = r'C:\Desktop\基础病理\research\_ruler.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])

colors = [RGBColor(0xDC, 0x26, 0x26), RGBColor(0x25, 0x63, 0xEB),
          RGBColor(0x0D, 0x94, 0x88), RGBColor(0xD9, 0x77, 0x06)]
for i, c in enumerate(colors):
    x = i * 3.3333
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(3.0), Inches(3.3333), Inches(1.5))
    sh.fill.solid()
    sh.fill.fore_color.rgb = c
    sh.line.fill.background()
    sh.text_frame.text = 'x=%.2f' % x
    sh.text_frame.paragraphs[0].runs[0].font.size = __import__('pptx.util', fromlist=['Pt']).Pt(18)

prs.save(P)
print('saved', P)
